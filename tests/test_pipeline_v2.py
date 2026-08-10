"""Unit tests for Program Pipeline v2 (Dolu Excel Eksik Sütun Doldurucu)."""

from __future__ import annotations

import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook

from src.program_pipeline.excel_writer_v2 import (
    is_cell_empty,
    format_field_value_v2,
    write_missing_fields_to_populated_excel,
)
from src.program_pipeline.pipeline_v2 import extract_urls_from_source_cell
from src.program_pipeline.models import ExtractedField, ProgramRecord


def test_is_cell_empty():
    assert is_cell_empty(None) is True
    assert is_cell_empty("") is True
    assert is_cell_empty("   ") is True
    assert is_cell_empty("—") is True
    assert is_cell_empty("[BOŞ]") is True
    assert is_cell_empty("English") is False
    assert is_cell_empty("2,25") is False


def test_format_field_value_v2():
    # 1. GPA formatting (Turkish comma decimal)
    gpa_field = ExtractedField(value=2.25, confidence=0.95)
    assert format_field_value_v2("Minimum GPA", gpa_field) == "2,25"

    gpa_str_field = ExtractedField(value="3.0", confidence=0.95)
    assert format_field_value_v2("Minimum GPA", gpa_str_field) == "3,0"

    # 2. Boolean formatting
    bool_field = ExtractedField(value=True, confidence=0.90)
    assert format_field_value_v2("Portfolyo Gerekli Mi", bool_field) == "Evet"

    # 3. String formatting
    dur_field = ExtractedField(value="60 units", confidence=0.90)
    assert format_field_value_v2("Eğitim Süresi", dur_field) == "60 units"


def test_extract_urls_from_source_cell():
    cell_text = (
        "Program: https://example.edu/programs/cs\n"
        "Kabul: https://example.edu/admissions\n"
        "Ücret: https://example.edu/tuition"
    )
    urls = extract_urls_from_source_cell(cell_text)
    assert urls["program_url"] == "https://example.edu/programs/cs"
    assert urls["admission_url"] == "https://example.edu/admissions"
    assert urls["tuition_url"] == "https://example.edu/tuition"

    raw_text = "https://example.edu/prog https://example.edu/adm"
    urls_raw = extract_urls_from_source_cell(raw_text)
    assert urls_raw["program_url"] == "https://example.edu/prog"
    assert urls_raw["admission_url"] == "https://example.edu/adm"


def test_write_missing_fields_to_populated_excel_preserves_existing_data():
    """Dolu olan hücrelerin üzerine yazılmadığını, sadece boş hücrelerin doldurulduğunu doğrula."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Programlar"

    # Başlık satırı
    headers = [
        "Üniversite ID", "Üniversite Adı", "Program Adı", "Fakülte",
        "Eğitim Süresi", "Minimum GPA", "Minimum IELTS", "Kabul Verisi Kaynağı"
    ]
    ws.append(headers)

    # Örnek dolu satır (Fakülte ve IELTS DOLU; Eğitim Süresi ve GPA BOŞ)
    ws.append([
        "UNI-001", "Test University", "Master of Science in Computer Science",
        "School of Engineering", "", "", "7.0", "Kabul: https://example.edu/admissions"
    ])

    with tempfile.TemporaryDirectory() as tmp_dir:
        input_path = Path(tmp_dir) / "populated_test.xlsx"
        output_path = Path(tmp_dir) / "output_test.xlsx"
        wb.save(input_path)
        wb.close()

        # Çıkarılan yeni veriler (Fakülte "New School", GPA 3.25, IELTS 6.5, Duration "2 yıl")
        record = ProgramRecord(
            university_name="Test University",
            program_name=ExtractedField(value="Master of Science in Computer Science", confidence=0.99),
            fields={
                "Fakülte": ExtractedField(value="New School of CS", confidence=0.95),  # Dolu hücre -> DEĞİŞTİRİLMEMELİ
                "Eğitim Süresi": ExtractedField(value="2 yıl / 30 units", confidence=0.95),  # Boş -> DOLDURULMALI
                "Minimum GPA": ExtractedField(value=3.25, confidence=0.95),  # Boş -> DOLDURULMALI ("3,25")
                "Minimum IELTS": ExtractedField(value="6.5", confidence=0.95),  # Dolu hücre ("7.0") -> DEĞİŞTİRİLMEMELİ
            }
        )

        col_map = {h: i + 1 for i, h in enumerate(headers)}

        stats = write_missing_fields_to_populated_excel(
            populated_excel_path=input_path,
            output_path=output_path,
            program_records=[record],
            column_map=col_map,
            sheet_name="Programlar",
            header_row=1,
            confidence_threshold=0.85,
        )

        assert stats["total_rows"] == 1
        assert stats["existing_cells_preserved"] >= 2  # Fakülte ve IELTS korundu
        assert stats["empty_cells_filled"] == 2  # Süre ve GPA dolduruldu

        # Çıktı dosyasını oku ve içerikleri kontrol et
        out_wb = load_workbook(output_path, data_only=True)
        out_ws = out_wb["Programlar"]

        # Fakülte (orijinal: "School of Engineering" -> DEĞİŞMEDİ)
        assert out_ws.cell(row=2, column=4).value == "School of Engineering"

        # Eğitim Süresi (orijinal boş -> dolduruldu: "2 yıl / 30 units")
        assert out_ws.cell(row=2, column=5).value == "2 yıl / 30 units"

        # Minimum GPA (orijinal boş -> dolduruldu: "3,25")
        assert out_ws.cell(row=2, column=6).value == "3,25"

        # Minimum IELTS (orijinal: "7.0" -> DEĞİŞMEDİ)
        assert out_ws.cell(row=2, column=7).value == "7.0"

        out_wb.close()
