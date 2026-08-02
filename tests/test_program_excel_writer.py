"""Programlar çıktısında kaynak ve tekilleştirme davranışı."""

from openpyxl import Workbook, load_workbook

from src.program_pipeline.excel_writer import write_programs_to_excel
from src.program_pipeline.models import ExtractedField, ProgramRecord, TemplateInfo


def test_program_source_comment_and_admission_source_are_distinct(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Programlar"
    headers = [
        "Üniversite ID", "Üniversite Adı", "Program Adı",
        "Kabul Verisi Kaynağı", "Minimum GPA",
    ]
    ws.append(headers)
    wb.save(template)
    wb.close()

    program_url = "https://example.edu/programs/master-of-divinity"
    admission_url = "https://example.edu/admissions/requirements"
    program = ProgramRecord(
        university_id="U-1",
        university_name="Example Seminary",
        program_name=ExtractedField(
            value="Master of Divinity",
            source_url=program_url,
            evidence="Master of Divinity",
            confidence=1.0,
        ),
        program_url=program_url,
        fields={
            "Kabul Verisi Kaynağı": ExtractedField(
                value=admission_url,
                source_url=admission_url,
                evidence="Admission requirements",
                confidence=0.99,
            ),
            "Minimum GPA": ExtractedField(
                value=2.5,
                source_url=admission_url,
                evidence="minimum GPA of 2.5",
                confidence=0.95,
            ),
        },
    )
    info = TemplateInfo(
        sheet_name="Programlar",
        header_row=1,
        column_map={header: index for index, header in enumerate(headers, 1)},
    )

    write_programs_to_excel([program], template, output, info)

    wb = load_workbook(output)
    ws = wb["Programlar"]
    assert ws["C2"].value == "Master of Divinity"
    assert program_url in ws["C2"].comment.text
    assert ws["D2"].value == admission_url
    assert ws["D2"].value != program_url
    assert ws["E2"].value == 2.5
    wb.close()
