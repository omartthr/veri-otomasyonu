"""Excel yazma testleri."""

from __future__ import annotations

import pytest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.excel_writer import (
    _get_available_output_path,
    _sanitize_cell_value,
    generate_output_filename,
    write_results,
)
from src.models import (
    CandidateResult, CandidateURL, DecisionType, ResearchCategory,
    RowStatus, ScoreBreakdown, UniversityResearchResult,
)


@pytest.fixture
def sample_excel(tmp_path: Path) -> Path:
    """Yazma testleri için örnek Excel."""
    wb = Workbook()
    ws = wb.active
    headers = [
        "ID", "Ad", "Şehir", "Tür", "Web", "F", "G", "H", "I",
        "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 15

    row2 = ["UNI-001", "Test Uni", "İst", "V", "https://testuni.edu",
            "Evet", "Evet", "Evet", "Evet",
            None, None, None, None, "Ali", "Bekliyor", None,
            "q1", "r1", "s1", "t1"]
    for col, val in enumerate(row2, 1):
        ws.cell(row=2, column=col, value=val)

    row3 = ["UNI-002", "Sample Uni", "Ank", "D", "https://sampleuni.edu.tr",
            "Hayır", "Evet", "Hayır", "Evet",
            "https://sampleuni.edu.tr/prog", None,
            "https://sampleuni.edu.tr/eng", None,
            None, None, None, "q2", "r2", None, None]
    for col, val in enumerate(row3, 1):
        ws.cell(row=3, column=col, value=val)

    fp = tmp_path / "write_test.xlsx"
    wb.save(fp)
    wb.close()
    return fp


class TestExcelWriter:
    def test_output_filename(self, tmp_path: Path):
        orig = tmp_path / "unis.xlsx"
        out = generate_output_filename(orig)
        assert out.name == "unis_doldurulmus.xlsx"
        assert out.suffix == ".xlsx"

    def test_qrst_preserved(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "out.xlsx"
        result = UniversityResearchResult(
            row_number=2, university_name="Test Uni", university_id="UNI-001",
            best_results={
                ResearchCategory.PROGRAM: CandidateResult(
                    candidate=CandidateURL(url="https://testuni.edu/prog", category=ResearchCategory.PROGRAM),
                    decision=DecisionType.ACCEPTED,
                ),
            },
            status=RowStatus.COMPLETED,
            notes=["HTTP 200; puan=95; uzun teknik günlük Excel'e yazılmamalı"],
        )
        write_results(sample_excel, out, [result])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=17).value == "q1"
        assert ws.cell(row=2, column=18).value == "r1"
        assert ws.cell(row=2, column=19).value == "s1"
        assert ws.cell(row=2, column=20).value == "t1"
        assert ws.cell(row=2, column=16).value is None
        wb.close()

    def test_accepted_written_but_rejected_skipped(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "out2.xlsx"
        result = UniversityResearchResult(
            row_number=2, university_name="Test Uni",
            best_results={
                ResearchCategory.PROGRAM: CandidateResult(
                    candidate=CandidateURL(url="https://testuni.edu/prog", category=ResearchCategory.PROGRAM),
                    decision=DecisionType.ACCEPTED,
                ),
                ResearchCategory.ADMISSION: CandidateResult(
                    candidate=CandidateURL(url="https://testuni.edu/adm", category=ResearchCategory.ADMISSION),
                    decision=DecisionType.REJECTED,
                ),
            }, status=RowStatus.MANUAL_REVIEW,
        )
        write_results(sample_excel, out, [result])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=10).value == "https://testuni.edu/prog"
        assert ws.cell(row=2, column=11).value is None
        assert ws.cell(row=2, column=14).value == "Ali"
        wb.close()

    def test_manual_review_url_cleared_and_status_selected(
        self, sample_excel: Path, tmp_path: Path
    ):
        out = tmp_path / "manual_review.xlsx"
        result = UniversityResearchResult(
            row_number=2,
            university_name="Test Uni",
            best_results={
                ResearchCategory.PROGRAM: CandidateResult(
                    candidate=CandidateURL(
                        url="https://testuni.edu/program-candidate",
                        category=ResearchCategory.PROGRAM,
                    ),
                    decision=DecisionType.MANUAL_REVIEW,
                    rejection_reason="Yıl bilgisi elle kontrol edilmeli",
                ),
            },
            status=RowStatus.MANUAL_REVIEW,
        )

        write_results(sample_excel, out, [result])

        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=10).value is None
        assert ws.cell(row=2, column=15).value == "Manuel Kontrol"
        wb.close()

    def test_unverified_existing_url_is_cleared(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "unverified_existing.xlsx"
        result = UniversityResearchResult(
            row_number=3,
            university_name="Sample Uni",
            best_results={ResearchCategory.PROGRAM: None},
            notes=[
                "Program listesi için doğrudan ve doğrulanabilir resmî kaynak bulunamadı."
            ],
        )
        write_results(sample_excel, out, [result])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=3, column=10).value is None
        assert ws.cell(row=3, column=12).value == "https://sampleuni.edu.tr/eng"
        assert ws.cell(row=3, column=15).value == "Manuel Kontrol"
        assert ws.cell(row=3, column=16).value == (
            "Program listesi için doğrudan ve doğrulanabilir resmî kaynak bulunamadı."
        )
        wb.close()

    def test_hayir_urls_preserved(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "out3.xlsx"
        result = UniversityResearchResult(
            row_number=3, university_name="Sample Uni",
            best_results={
                ResearchCategory.ADMISSION: CandidateResult(
                    candidate=CandidateURL(url="https://sampleuni.edu.tr/adm", category=ResearchCategory.ADMISSION),
                    decision=DecisionType.ACCEPTED,
                ),
            }, status=RowStatus.COMPLETED,
        )
        write_results(sample_excel, out, [result])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=3, column=10).value == "https://sampleuni.edu.tr/prog"
        assert ws.cell(row=3, column=11).value == "https://sampleuni.edu.tr/adm"
        assert ws.cell(row=3, column=12).value == "https://sampleuni.edu.tr/eng"
        wb.close()

    def test_cannot_overwrite_original(self, sample_excel: Path):
        with pytest.raises(ValueError):
            write_results(sample_excel, sample_excel, [])

    def test_status_updated(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "out4.xlsx"
        result = UniversityResearchResult(
            row_number=2, university_name="Test Uni",
            best_results={}, status=RowStatus.NOT_FOUND,
            notes=["Kaynak bulunamadı"],
        )
        write_results(sample_excel, out, [result])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=2, column=15).value == "Web sitesi yok/bulunamadı"
        wb.close()

    def test_formatting_preserved(self, sample_excel: Path, tmp_path: Path):
        out = tmp_path / "out5.xlsx"
        write_results(sample_excel, out, [])
        wb = load_workbook(out)
        ws = wb.active
        assert ws.cell(row=1, column=1).font.bold is True
        assert ws.column_dimensions["A"].width == 15
        wb.close()

    def test_sanitize_formula_injection(self):
        assert _sanitize_cell_value("=SUM(A1)") == "'=SUM(A1)"
        assert _sanitize_cell_value("+cmd('c')") == "'+cmd('c')"
        assert _sanitize_cell_value("Normal") == "Normal"
        assert _sanitize_cell_value("https://x.com") == "https://x.com"
        assert _sanitize_cell_value("") == ""

    def test_existing_output_gets_numbered_path(self, tmp_path: Path):
        target = tmp_path / "report_doldurulmus.xlsx"
        target.write_bytes(b"existing")
        available = _get_available_output_path(target)
        assert available.name == "report_doldurulmus_2.xlsx"

    def test_write_does_not_overwrite_existing_output(self, sample_excel: Path, tmp_path: Path):
        target = tmp_path / "out.xlsx"
        target.write_bytes(b"keep-me")
        saved = write_results(sample_excel, target, [])
        assert saved.name == "out_2.xlsx"
        assert target.read_bytes() == b"keep-me"
        assert saved.exists()
