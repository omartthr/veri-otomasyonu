"""Phoenix-benzeri bir kurum için ağsız tam kalite pilotu."""

import asyncio

import pytest
from openpyxl import Workbook, load_workbook

from src.audit.evidence_store import EvidenceStore
from src.audit.review_queue import ReviewQueue
from src.program_pipeline.content_sections import extract_html_sections
from src.program_pipeline.excel_writer import write_programs_to_excel
from src.program_pipeline.general_policy_extractor import GeneralPolicyExtractor
from src.program_pipeline.models import SourceUniversityRow, TemplateInfo
from src.program_pipeline.pipeline import ProgramPipeline


PROGRAM_URL = "https://ps.edu/master-of-divinity"
ADMISSION_URL = "https://ps.edu/admissions/"
TUITION_URL = "https://ps.edu/tuition-and-fees/"

PROGRAM_HTML = b"""
<html><head><title>Master of Divinity | Phoenix Seminary</title></head>
<body><main>
<h1>Master of Divinity</h1>
<p>The Master of Divinity is a 92 credit hour degree that can be completed in 3 years.</p>
<h2>Program Requirements</h2>
<p>The curriculum prepares graduates for careers in pastoral ministry, chaplaincy, teaching, and church leadership.</p>
<p>An integrated practical theology internship opportunity is included in the program.</p>
</main></body></html>
"""

ADMISSION_HTML = b"""
<html><head><title>Admissions | Phoenix Seminary</title></head>
<body><main>
<h1>Admission Requirements</h1>
<p>Applicants must complete an online application and submit a non-refundable application fee of $40.</p>
<p>Applicants must submit official transcripts from all previously attended institutions, a written personal testimony essay, and completed recommendation forms.</p>
<p>A bachelor's degree from an accredited institution and a minimum GPA of 2.5 are required.</p>
<p>If English is not the applicant's first language, the applicant must demonstrate English proficiency with a minimum TOEFL iBT score of 92 or IELTS score of 6.5. Test results must be no more than two years old.</p>
<p>The Admission Committee may request a personal interview and may require a credential evaluation for foreign degrees.</p>
<p>Applications are accepted no more than one year before the intended start date.</p>
</main></body></html>
"""

TUITION_HTML = b"""
<html><head><title>Tuition and Fees | Phoenix Seminary</title></head>
<body><main>
<h1>Tuition and Fees</h1>
<p>For academic year 2026-2027, Master of Divinity tuition is $580 per credit hour.</p>
</main></body></html>
"""

NOISE_HTML = b"""
<html><head><title>Register | Phoenix Seminary</title></head>
<body><main><h1>Register</h1><p>Register for an upcoming public event.</p></main></body></html>
"""


def _merge_general(html: bytes, url: str, url_type: str, target: dict, page_texts: dict):
    fields = GeneralPolicyExtractor().extract_from_html(html, url, url_type)
    for name, field in fields.items():
        if name not in target or field.confidence > target[name].confidence:
            target[name] = field
    page_texts[url] = " ".join(
        f"{section.heading or ''} {section.text}"
        for section in extract_html_sections(html, url)
    )


@pytest.mark.asyncio
async def test_pilot_rejects_noise_and_writes_rich_evidence_backed_row(tmp_path):
    pipeline = ProgramPipeline(confidence_threshold=0.85, resume=False)
    pages = {
        PROGRAM_URL: PROGRAM_HTML,
        "https://ps.edu/register/": NOISE_HTML,
    }

    async def fake_fetch(url):
        return pages.get(url)

    pipeline._fetch_html = fake_fetch

    general_fields = {}
    page_texts = {}
    _merge_general(ADMISSION_HTML, ADMISSION_URL, "admission", general_fields, page_texts)
    _merge_general(TUITION_HTML, TUITION_URL, "tuition", general_fields, page_texts)

    row = SourceUniversityRow(
        row_number=2,
        university_id="SCORECARD-381459",
        university_name="Phoenix Seminary",
        official_website="https://ps.edu/",
        program_url="https://ps.edu/degree-programs/",
        admission_url=ADMISSION_URL,
        language_url=ADMISSION_URL,
        tuition_url=TUITION_URL,
    )
    audit = EvidenceStore(tmp_path / "audit")
    queue = ReviewQueue(tmp_path / "review.jsonl")
    semaphore = asyncio.Semaphore(1)

    program, _, _ = await pipeline._process_program(
        prog_url=PROGRAM_URL,
        row=row,
        official_domain="ps.edu",
        general_fields=general_fields,
        page_texts=page_texts,
        catalog_sections=[],
        template_info=TemplateInfo(),
        audit=audit,
        review_q=queue,
        semaphore=semaphore,
    )
    noise, _, _ = await pipeline._process_program(
        prog_url="https://ps.edu/register/",
        row=row,
        official_domain="ps.edu",
        general_fields=general_fields,
        page_texts=page_texts,
        catalog_sections=[],
        template_info=TemplateInfo(),
        audit=audit,
        review_q=queue,
        semaphore=semaphore,
    )
    audit.close()

    assert noise is None
    assert program is not None
    assert program.program_name.value == "Master of Divinity"
    assert sum(field.value is not None for field in program.fields.values()) >= 15
    assert program.get_field("Minimum GPA").value == 2.5
    assert program.get_field("Minimum IELTS").value == 6.5
    assert program.get_field("Minimum TOEFL").value == 92
    assert program.get_field("Başvuru Ücreti Gerekli Mi").value is True
    assert str(program.get_field("Mülakat Gerekli Mi").value).startswith("Koşullu:")

    headers = [
        "Üniversite ID", "Üniversite Adı", "Program Adı", "Fakülte",
        "Eğitim Süresi", "Eğitim Dili", "Program Ranking", "Yıllık Ücret",
        "Staj İmkanı", "Co-op İmkanı", "Exchange İmkanı", "Çift Diploma",
        "Mezun Maaşı", "Mezun İstihdamı", "Meslek Çıktıları", "Yök Denklik",
        "Minimum GPA", "Minimum IB", "Minimum SAT", "Minimum ACT",
        "Minimum IELTS", "Minimum TOEFL", "Minimum Duolingo", "AP Gereksinimi",
        "A-Level Gereksinimi", "Portfolyo Gerekli Mi", "Referans Sayısı",
        "Essay Gerekli Mi", "Mülakat Gerekli Mi", "CV Gerekli Mi",
        "Video Başvuru Gerekli Mi", "Kabul Verisi Kaynağı", "Başvuru Türü",
        "Başvuru Dönemi", "Akademik Yıl", "Başlangıç Tarihi",
        "Son Başvuru Tarihi", "Son Başvuru Metni", "Başvuru Ücreti",
        "Başvuru Ücreti Gerekli Mi", "Offline Formlar", "Başvuru Belgeleri",
        "Zorunlu Başvuru Maddeleri", "İnceleme Gerektiren Maddeler",
        "İngilizce Yeterlilik Gerekli Mi", "Standart Test Politikası",
        "Transkript/GED Gereksinimi", "CommonApp Checklist Uyarıları",
    ]
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Programlar"
    ws.append(headers)
    wb.save(template)
    wb.close()

    info = TemplateInfo(
        sheet_name="Programlar",
        header_row=1,
        column_map={header: index for index, header in enumerate(headers, 1)},
    )
    write_programs_to_excel([program], template, output, info)

    wb = load_workbook(output)
    ws = wb["Programlar"]
    filled = sum(ws.cell(2, column).value not in (None, "") for column in range(1, 49))
    assert filled >= 18
    assert ws["C2"].comment and PROGRAM_URL in ws["C2"].comment.text
    assert ws["AF2"].value == ADMISSION_URL
    wb.close()
