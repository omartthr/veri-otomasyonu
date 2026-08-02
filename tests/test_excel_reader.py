"""Excel okuma testleri."""

from __future__ import annotations

import pytest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.excel_reader import read_excel, count_missing_fields, get_universities_needing_research
from src.models import ResearchCategory


@pytest.fixture
def sample_excel(tmp_path: Path) -> Path:
    """Örnek Excel fixture."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Üniversiteler"

    headers = [
        "Üniversite ID", "Üniversite Adı", "Şehir", "Kurum Türü", "Web Sitesi",
        "Program Listesi Eksik?", "Kabul Şartları Eksik?", "Dil Şartları Eksik?",
        "Ücret Bilgisi Eksik?", "Program Listesi URL", "Kabul Şartları URL",
        "Dil Şartları URL", "Ücret URL", "Atanan Kişi", "Durum", "Notlar",
        "Çekim 1", "Çekim 2", "Çekim 3", "Çekim 4",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    # Satır 2 — Tüm alanlar eksik
    row2 = ["UNI-001", "Test University", "İstanbul", "Vakıf",
            "https://www.testuni.edu", "Evet", "Evet", "Evet", "Evet",
            None, None, None, None, "Ali Yılmaz", "Bekliyor", None,
            "çekim1", "çekim2", "çekim3", "çekim4"]
    for col, val in enumerate(row2, 1):
        ws.cell(row=2, column=col, value=val)

    # Satır 3 — Bazıları Hayır
    row3 = ["UNI-002", "Sample University", "Ankara", "Devlet",
            "https://www.sampleuni.edu.tr", "Hayır", "Evet", "Hayır", "Evet",
            "https://www.sampleuni.edu.tr/programs", None,
            "https://www.sampleuni.edu.tr/english", None,
            None, None, None, "q", "r", None, None]
    for col, val in enumerate(row3, 1):
        ws.cell(row=3, column=col, value=val)

    # Satır 4 — Hiç eksik yok
    row4 = ["UNI-003", "Complete University", "İzmir", "Devlet",
            "https://completeuni.edu", "Hayır", "Hayır", "Hayır", "Hayır",
            "https://completeuni.edu/programs", "https://completeuni.edu/admissions",
            "https://completeuni.edu/english", "https://completeuni.edu/tuition",
            None, None, None, None, None, None, None]
    for col, val in enumerate(row4, 1):
        ws.cell(row=4, column=col, value=val)

    # Satır 5 — gerekli değil
    row5 = ["UNI-004", "Partial University", "Bursa", "Vakıf",
            "https://partialuni.edu", "Evet", "gerekli değil", "Evet",
            "gerekli değil", None, None, None, None, None, None, None,
            None, None, None, None]
    for col, val in enumerate(row5, 1):
        ws.cell(row=5, column=col, value=val)

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 30

    file_path = tmp_path / "test_universities.xlsx"
    wb.save(file_path)
    wb.close()
    return file_path


class TestExcelReader:
    def test_read_returns_correct_count(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        assert len(rows) == 4

    def test_parses_fields(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        first = rows[0]
        assert first.university_id == "UNI-001"
        assert first.university_name == "Test University"
        assert first.city == "İstanbul"
        assert first.website == "https://www.testuni.edu"
        assert first.program_missing == "Evet"

    def test_row_numbers(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        assert rows[0].row_number == 2
        assert rows[3].row_number == 5

    def test_all_missing(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        missing = rows[0].get_missing_categories()
        assert len(missing) == 4

    def test_partial_missing(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        missing = rows[1].get_missing_categories()
        assert len(missing) == 2
        assert ResearchCategory.ADMISSION in missing
        assert ResearchCategory.TUITION in missing
        assert ResearchCategory.PROGRAM not in missing

    def test_none_missing(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        assert len(rows[2].get_missing_categories()) == 0

    def test_gerekli_degil_ignored(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        missing = rows[3].get_missing_categories()
        assert len(missing) == 2
        assert ResearchCategory.PROGRAM in missing
        assert ResearchCategory.LANGUAGE in missing

    def test_count_missing_fields(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        stats = count_missing_fields(rows)
        assert stats["total"] == 8

    def test_needing_research(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        needing = get_universities_needing_research(rows)
        assert len(needing) == 3
        names = [r.university_name for r in needing]
        assert "Complete University" not in names

    def test_file_not_found(self, tmp_path: Path):
        with pytest.raises(FileNotFoundError):
            read_excel(tmp_path / "no.xlsx")

    def test_invalid_extension(self, tmp_path: Path):
        f = tmp_path / "test.csv"
        f.write_text("x")
        with pytest.raises(ValueError):
            read_excel(f)

    def test_existing_urls(self, sample_excel: Path):
        rows = read_excel(sample_excel)
        second = rows[1]
        assert second.program_url == "https://www.sampleuni.edu.tr/programs"
        assert second.admission_url is None
