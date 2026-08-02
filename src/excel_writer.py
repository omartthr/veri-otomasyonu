"""Excel dosyasına araştırma sonuçlarını yazma modülü.

Orijinal dosyanın biçimini (formüller, stiller, renkler, sütun genişlikleri,
filtreler, veri doğrulamaları) koruyarak yeni bir dosya üretir.
Orijinal dosyanın üzerine asla yazmaz.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.models import (
    DecisionType,
    ResearchCategory,
    UniversityResearchResult,
    URL_RESULT_COLUMNS,
    STATUS_COLUMN,
    NOTES_COLUMN,
    PROTECTED_COLUMNS,
)

logger = logging.getLogger(__name__)

_CATEGORY_NOTE_NAMES: dict[ResearchCategory, str] = {
    ResearchCategory.PROGRAM: "Program listesi",
    ResearchCategory.ADMISSION: "Kabul şartları",
    ResearchCategory.LANGUAGE: "Dil şartları",
    ResearchCategory.TUITION: "Ücret bilgisi",
}

# Sütun harfi → numara çevirici
_COL_LETTER_TO_NUM: dict[str, int] = {chr(i): i - 64 for i in range(65, 91)}  # A=1 … Z=26


def _col_num(letter: str) -> int:
    """Sütun harfini numaraya çevir (A=1)."""
    return _COL_LETTER_TO_NUM[letter.upper()]


# ─────────── Excel formül enjeksiyonuna karşı koruma ───────────

_FORMULA_CHARS = {"=", "+", "-", "@", "\t", "\r", "\n"}


def _sanitize_cell_value(value: str) -> str:
    """Excel formül enjeksiyonuna karşı metin değeri güvenli hale getir.

    URL'ler için değişiklik yapmaz (http/https ile başlayan).
    Diğer metinlerde tehlikeli başlangıç karakterlerini temizler.
    """
    if not value:
        return value

    # URL'leri koru
    if value.strip().lower().startswith(("http://", "https://")):
        return value

    # Tehlikeli başlangıç karakteri varsa önek ekle
    if value[0] in _FORMULA_CHARS:
        return "'" + value

    return value


def generate_output_filename(original_path: Path) -> Path:
    """Çıktı dosya adını oluştur.

    Format: orijinal-ad_doldurulmus.xlsx
    """
    stem = original_path.stem
    new_name = f"{stem}_doldurulmus.xlsx"
    return original_path.parent / new_name


def _get_available_output_path(output_path: Path) -> Path:
    """Mevcut/kilitli çıktının üzerine yazmadan uygun numaralı yol üret."""
    if not output_path.exists():
        return output_path
    for index in range(2, 10_000):
        candidate = output_path.with_name(
            f"{output_path.stem}_{index}{output_path.suffix}"
        )
        if not candidate.exists():
            return candidate
    raise FileExistsError(f"Uygun çıktı dosya adı üretilemedi: {output_path}")


def write_results(
    original_path: Path | str,
    output_path: Optional[Path | str],
    results: list[UniversityResearchResult],
    sheet_name: Optional[str] = None,
) -> Path:
    """Araştırma sonuçlarını orijinal Excel dosyasının bir kopyasına yaz.

    Kurallar:
    - Yalnızca tüm doğrulamaları geçen ACCEPTED sonuçlar URL sütunlarına yazılır.
    - MANUAL_REVIEW/REJECTED/NOT_FOUND sonuçlarında araştırılan URL hücresi boş bırakılır.
    - Doğrulanamayan alan bulunan satırın durumu "Manuel Kontrol" olur.
    - Q-T sütunlarına dokunulmaz.
    - "Hayır" veya "gerekli değil" olan alanlara dokunulmaz.
    - Formüller, stiller, renkler, sütun genişlikleri korunur.
    - Orijinal dosyanın üzerine yazılmaz.

    Args:
        original_path: Orijinal .xlsx dosya yolu
        output_path: Çıktı dosya yolu (None ise otomatik oluşturulur)
        results: Üniversite araştırma sonuçları listesi
        sheet_name: Yazılacak sayfa adı (None = aktif sayfa)

    Returns:
        Çıktı dosyasının yolu
    """
    original_path = Path(original_path)
    if not original_path.exists():
        raise FileNotFoundError(f"Orijinal dosya bulunamadı: {original_path}")

    # Çıktı yolunu belirle
    if output_path is None:
        output_path = generate_output_filename(original_path)
    else:
        output_path = Path(output_path)

    # Orijinalin üzerine yazmayı engelle
    if output_path.resolve() == original_path.resolve():
        raise ValueError("Çıktı dosyası orijinal dosyanın üzerine yazılamaz!")

    # Önceki çıktı açık, OneDrive tarafından kilitli veya korunuyor olabilir.
    # Hiçbir mevcut çıktının üzerine yazma; otomatik _2, _3 ... adı seç.
    output_path = _get_available_output_path(output_path)

    logger.info("Orijinal dosya kopyalanıyor: %s → %s", original_path, output_path)

    # Orijinal dosyayı formüller dahil yükle (data_only=False)
    wb = load_workbook(original_path, data_only=False)
    if sheet_name:
        ws = wb[sheet_name]
    elif "Takip" in wb.sheetnames:
        ws = wb["Takip"]
    else:
        ws = wb.active

    if ws is None:
        raise ValueError("Excel dosyasında aktif sayfa bulunamadı.")

    selected_sheet_name = ws.title
    protected_snapshot = {
        (row, _col_num(col)): ws.cell(row=row, column=_col_num(col)).value
        for row in range(1, ws.max_row + 1)
        for col in PROTECTED_COLUMNS
    }

    # Sonuçları satır numarasına göre indexle
    results_by_row: dict[int, UniversityResearchResult] = {
        r.row_number: r for r in results
    }

    for row_num, research_result in results_by_row.items():
        _write_row_results(ws, row_num, research_result)

    # Çıktı dizinini oluştur
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Kaydet
    wb.save(output_path)
    wb.close()

    # Çıktıyı yeniden aç; dosyanın okunabildiğini ve Q-T'nin değişmediğini doğrula.
    check_wb = load_workbook(output_path, data_only=False)
    check_ws = check_wb[selected_sheet_name]
    for (row, col), expected in protected_snapshot.items():
        actual = check_ws.cell(row=row, column=col).value
        if actual != expected:
            check_wb.close()
            raise ValueError(f"Korunan hücre değişti: {check_ws.cell(row=row, column=col).coordinate}")
    check_wb.close()

    logger.info("Sonuçlar yazıldı: %s", output_path)
    return output_path


def _write_row_results(
    ws: Worksheet,
    row_num: int,
    result: UniversityResearchResult,
) -> None:
    """Tek bir satıra araştırma sonuçlarını yaz."""

    for category, col_letter in URL_RESULT_COLUMNS.items():
        col_num = _col_num(col_letter)

        # Bu kategori için en iyi sonucu al
        best = result.best_results.get(category)

        if category not in result.best_results:
            # Araştırılması istenmeyen kategoriye dokunma.
            continue

        if best is None or best.decision != DecisionType.ACCEPTED:
            # Eski/geçersiz bir bağlantı mevcut olsa bile doğrulanamadıysa temizle.
            ws.cell(row=row_num, column=col_num).value = None
            continue

        # URL'yi yaz
        url = best.candidate.url
        ws.cell(row=row_num, column=col_num, value=url)

    # Durum sütunu (O)
    status_col = _col_num(STATUS_COLUMN)
    new_status = result.determine_status()
    ws.cell(row=row_num, column=status_col, value=new_status.value)

    # Notlar sütunu (P) — karar verilerinden yeniden üret; eski oturumlardaki
    # teknik günlüklerin veya uzun hata metinlerinin Excel'e sızmasına izin verme.
    notes_col = _col_num(NOTES_COLUMN)
    unresolved = [
        category
        for category, best in result.best_results.items()
        if best is None or best.decision != DecisionType.ACCEPTED
    ]
    concise_note = _missing_source_note(unresolved)
    notes_cell = ws.cell(row=row_num, column=notes_col)
    notes_cell.value = _sanitize_cell_value(concise_note) if concise_note else None


def _missing_source_note(categories: list[ResearchCategory]) -> str:
    """Doğrulanamayan kategoriler için tek kısa Excel notu üret."""
    if not categories:
        return ""
    names = [_CATEGORY_NOTE_NAMES[category] for category in categories]
    subject = names[0] if len(names) == 1 else ", ".join(names[:-1]) + f" ve {names[-1]}"
    return f"{subject} için doğrudan ve doğrulanabilir resmî kaynak bulunamadı."
