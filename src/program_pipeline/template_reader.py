"""Template Excel okuyucu — Programlar sheet analizi.

Mevcut workbook'tan:
- Programlar sheet'ini bulur
- Başlık satırını otomatik tespit eder
- 48 sütunun haritasını çıkarır
- Mevcut program satırlarını okur (duplicate kontrolü için)
- Programlar BOŞSA → Universiteler sheet'inden üniversite adlarını alır
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from src.program_pipeline.models import TemplateInfo

logger = logging.getLogger(__name__)

# Alias tablosu
COLUMN_ALIASES: dict[str, list[str]] = {
    "Üniversite ID": ["Üniversite ID", "University ID", "Uni ID"],
    "Üniversite Adı": ["Üniversite Adı", "University Name", "Üniversite"],
    "Program Adı": ["Program Adı", "Program Name", "Program"],
    "Minimum GPA": ["Minimum GPA", "GPA", "Min GPA"],
    "Minimum IELTS": ["Minimum IELTS", "IELTS", "Min IELTS"],
    "Minimum TOEFL": ["Minimum TOEFL", "TOEFL", "Min TOEFL"],
    "Minimum Duolingo": ["Minimum Duolingo", "Duolingo", "Min Duolingo"],
    "Yıllık Ücret": ["Yıllık Ücret", "Annual Tuition", "Tuition", "Ücret"],
    "Başvuru Ücreti": ["Başvuru Ücreti", "Application Fee", "App Fee"],
}


def _normalize_header(text: str) -> str:
    return str(text).strip().lower().replace("  ", " ")


def _find_programlar_sheet(wb) -> str:
    """Programlar sheet'ini bul."""
    candidates = ["Programlar", "Programs", "PROGRAMLAR"]
    for c in candidates:
        if c in wb.sheetnames:
            return c
    return wb.sheetnames[0] if wb.sheetnames else ""


def _find_universiteler_sheet(wb) -> Optional[str]:
    """Universiteler sheet'ini bul."""
    candidates = ["Universiteler", "Universities", "Üniversiteler", "UNIVERSITELER"]
    for c in candidates:
        if c in wb.sheetnames:
            return c
    return None


def _find_header_row(ws, max_scan: int = 5) -> int:
    """İlk 5 satırda en çok dolu sütun içeren satırı başlık satırı kabul et."""
    best_row, best_count = 1, 0
    for row_idx in range(1, min(max_scan + 1, (ws.max_row or 1) + 1)):
        filled = sum(
            1 for c in range(1, min((ws.max_column or 1), 60) + 1)
            if ws.cell(row=row_idx, column=c).value
        )
        if filled > best_count:
            best_count = filled
            best_row = row_idx
    return best_row


def _read_university_names_from_sheet(ws, header_row: int) -> list[str]:
    """Herhangi bir sheetten üniversite adlarını oku."""
    max_col = ws.max_column or 1
    headers = [
        str(ws.cell(row=header_row, column=c).value or "").strip()
        for c in range(1, max_col + 1)
    ]

    # "Üniversite Adı" veya "University Name" sütununu bul
    name_col = None
    for i, h in enumerate(headers):
        h_low = h.lower()
        if any(alias in h_low for alias in ["üniversite adı", "üniversite ad", "university name"]):
            name_col = i + 1
            break

    # Bulunamazsa 2. sütunu dene
    if name_col is None:
        name_col = 2

    names: list[str] = []
    for row_idx in range(header_row + 1, (ws.max_row or 1) + 1):
        val = ws.cell(row=row_idx, column=name_col).value
        if val:
            name = str(val).strip()
            if name and name not in names:
                names.append(name)
    return names


def read_template_excel(
    file_path: Path | str,
    sheet_name: Optional[str] = None,
) -> TemplateInfo:
    """Template Excel'i okuyarak TemplateInfo döndür.

    Programlar sheet'i boşsa Universiteler sheet'inden üniversite adlarını alır.

    Args:
        file_path: .xlsx şablon dosya yolu
        sheet_name: Programlar sheet adı (None = otomatik tespit)

    Returns:
        TemplateInfo nesnesi
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Template Excel bulunamadı: {file_path}")

    logger.info("Template Excel okunuyor: %s", file_path)
    wb = load_workbook(file_path, data_only=True)

    actual_sheet = sheet_name or _find_programlar_sheet(wb)
    if actual_sheet not in wb.sheetnames:
        raise ValueError(f"Sheet '{actual_sheet}' bulunamadı. Mevcut: {wb.sheetnames}")

    ws = wb[actual_sheet]
    header_row = _find_header_row(ws)

    # Sütun haritası oluştur (başlık adı → sütun numarası)
    column_map: dict[str, int] = {}
    for c in range(1, (ws.max_column or 1) + 1):
        cell_val = ws.cell(row=header_row, column=c).value
        if not cell_val:
            continue
        cell_str = str(cell_val).strip()
        column_map[cell_str] = c
        # Alias eşleştirmesi
        for canonical, aliases in COLUMN_ALIASES.items():
            if any(_normalize_header(a) == _normalize_header(cell_str) for a in aliases):
                if canonical not in column_map:
                    column_map[canonical] = c

    logger.info(
        "Programlar sheet: '%s', başlık satırı=%d, sütun sayısı=%d",
        actual_sheet, header_row, len(column_map),
    )

    # Mevcut program satırlarını oku
    university_names: list[str] = []
    existing_programs: list[tuple[str, str]] = []

    uni_col = column_map.get("Üniversite Adı", column_map.get("University Name"))
    prog_col = column_map.get("Program Adı", column_map.get("Program Name", column_map.get("Program")))

    if ws.max_row > header_row and uni_col:
        for row_idx in range(header_row + 1, ws.max_row + 1):
            uni_val = ws.cell(row=row_idx, column=uni_col).value
            prog_val = ws.cell(row=row_idx, column=prog_col).value if prog_col else None
            if uni_val:
                uni_str = str(uni_val).strip()
                if uni_str and uni_str not in university_names:
                    university_names.append(uni_str)
                if prog_val:
                    existing_programs.append((uni_str.lower(), str(prog_val).strip().lower()))

    # Programlar sheet'i boşsa → Universiteler sheet'inden al
    if not university_names:
        logger.info("Programlar sheet'i boş → Universiteler sheet'inden üniversite adları aranıyor...")
        uni_sheet_name = _find_universiteler_sheet(wb)
        if uni_sheet_name and uni_sheet_name in wb.sheetnames:
            ws_uni = wb[uni_sheet_name]
            uni_hdr = _find_header_row(ws_uni)
            university_names = _read_university_names_from_sheet(ws_uni, uni_hdr)
            logger.info("'%s' sheet'inden %d üniversite adı okundu.", uni_sheet_name, len(university_names))
        else:
            logger.warning("Universiteler sheet bulunamadı!")

    wb.close()
    return TemplateInfo(
        sheet_name=actual_sheet,
        header_row=header_row,
        column_map=column_map,
        university_names=university_names,
        existing_programs=existing_programs,
    )
