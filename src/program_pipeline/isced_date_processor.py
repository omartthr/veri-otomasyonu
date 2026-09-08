"""Fakülte ISCED & Tarih Düzeltici İşleyici Modülü.

Görevi:
1. Yüklenen Excel dosyasındaki 'Fakülte' sütununu UNESCO ISCED-F 2013 kodu karşılığına çevirir.
   - Eğer Fakülte hücresi boşsa veya net değilse, resmî domain üzerinden kural tabanlı hedefli tarama yaparak fakülte/bölüm bilgisini çeker ve ISCED kodunu yazar.
2. 'Son Başvuru Tarihi' sütunundaki tüm tarihleri YYYY-MM-DD formatına dönüştürür.
3. Diğer hiçbir sütun veya hücredeki veriye Dokunmaz (Zero Overwrite Guarantee).
"""

from __future__ import annotations

import asyncio
import logging
import shutil

from pathlib import Path
from typing import Callable, Optional, Any
from openpyxl import load_workbook

from src.config import get_settings
from src.domain_utils import extract_domain
from src.extractors.date_formatter import normalize_date_to_iso
from src.extractors.isced_mapper import map_to_isced_code
from src.program_pipeline.excel_writer_v2 import is_cell_empty
from src.program_pipeline.template_reader import read_template_excel
from src.program_pipeline.targeted_search import run_targeted_search_for_program
from src.fetch.page_fetcher import PageFetcher

logger = logging.getLogger(__name__)


async def process_isced_and_dates(
    input_excel_path: Path | str,
    output_excel_path: Path | str,
    progress_callback: Optional[Callable[[dict], None]] = None,
) -> dict[str, Any]:
    """Excel dosyasındaki Fakülte (ISCED) ve Son Başvuru Tarihi (YYYY-MM-DD) alanlarını günceller.

    Args:
        input_excel_path: Yüklenen Excel dosyası
        output_excel_path: Çıktı dosya yolu
        progress_callback: İlerleme bildirimi fonksiyonu

    Returns:
        İstatistik özet dict
    """
    input_excel_path = Path(input_excel_path)
    output_excel_path = Path(output_excel_path)
    output_excel_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Orijinal dosyayı çıktı yoluna kopyala (stil/renk kaybını önlemek için)
    shutil.copy2(input_excel_path, output_excel_path)

    template_info = read_template_excel(output_excel_path)
    wb = load_workbook(output_excel_path)
    ws = wb[template_info.sheet_name]
    header_row = template_info.header_row
    col_map = template_info.column_map

    # Sütun indekslerini tespit et
    fac_col = col_map.get("Fakülte") or col_map.get("Faculty") or col_map.get("Fakulte")
    date_col = col_map.get("Son Başvuru Tarihi") or col_map.get("Application Deadline") or col_map.get("Deadline")
    prog_col = col_map.get("Program Adı") or col_map.get("Program Name") or 3
    uni_col = col_map.get("Üniversite Adı") or col_map.get("University Name") or 2
    source_col = col_map.get("Kabul Verisi Kaynağı")

    if not fac_col and not date_col:
        wb.close()
        raise ValueError("Excel dosyasında 'Fakülte' veya 'Son Başvuru Tarihi' sütunu bulunamadı.")

    max_r = ws.max_row or header_row
    total_rows = max_r - header_row

    isced_updated_count = 0
    dates_updated_count = 0
    web_crawled_count = 0

    fetcher = PageFetcher()

    for idx, row_idx in enumerate(range(header_row + 1, max_r + 1)):
        uni_val = ws.cell(row=row_idx, column=uni_col).value if uni_col else ""
        prog_val = ws.cell(row=row_idx, column=prog_col).value if prog_col else ""
        fac_val = ws.cell(row=row_idx, column=fac_col).value if fac_col else ""
        date_val = ws.cell(row=row_idx, column=date_col).value if date_col else ""
        source_val = ws.cell(row=row_idx, column=source_col).value if source_col else ""

        uni_name = str(uni_val or "").strip()
        prog_name = str(prog_val or "").strip()
        fac_text = str(fac_val or "").strip()
        date_text = str(date_val or "").strip()

        if progress_callback:
            progress_callback({
                "current": idx + 1,
                "total": total_rows,
                "uni_name": uni_name,
                "prog_name": prog_name,
            })

        # ── 1. GÖREV: Son Başvuru Tarihi YYYY-MM-DD Biçimlendirme ──
        if date_col and date_text and not is_cell_empty(date_text):
            norm_date = normalize_date_to_iso(date_text)
            if norm_date != date_text:
                ws.cell(row=row_idx, column=date_col).value = norm_date
                dates_updated_count += 1

        # ── 2. GÖREV: Fakülte Sütununa ISCED Kodu Yazma ──
        if fac_col:
            # Mevcut veri zaten geçerli bir ISCED kodu mu? (Örn: "0613 - ...")
            if fac_text and fac_text.startswith(("0", "1")) and " - " in fac_text:
                continue

            # ISCED kodunu çıkar
            isced_code = map_to_isced_code(program_name=prog_name, faculty_name=fac_text)

            # Eğer doğrudan eşleşmediyse (0000 döndüyse) veya fakülte hücre boşsa -> Resmî Web Taraması Yap
            if (isced_code.startswith("0000") or is_cell_empty(fac_text)) and uni_name and prog_name:
                web_crawled_count += 1
                try:
                    # Kaynak hücresindeki URL veya domain üzerinden arama yap
                    source_url = ""
                    if source_val and "http" in str(source_val):
                        import re
                        m = re.search(r"https?://[^\s\n,]+", str(source_val))
                        if m:
                            source_url = m.group(0)

                    domain = extract_domain(source_url or f"{uni_name}.edu")
                    if domain:
                        discovered_urls = await run_targeted_search_for_program(
                            official_domain=domain,
                            program_name=prog_name,
                            missing_aspect="department faculty curriculum degree",
                        )
                        if discovered_urls:
                            fetched = await fetcher.fetch(discovered_urls[0])
                            if fetched and fetched.text_content:
                                # Sayfa içeriği üzerinden tekrar ISCED eşlemesi yap
                                isced_code = map_to_isced_code(
                                    program_name=prog_name,
                                    faculty_name=fac_text,
                                    text_content=fetched.text_content[:3000],
                                )
                except Exception as e:
                    logger.debug("ISCED web tarama hatası (%s): %s", prog_name, e)

            # Belirlenen ISCED kodunu Fakülte sütununa yaz
            if isced_code and not isced_code.startswith("0000"):
                ws.cell(row=row_idx, column=fac_col).value = isced_code
                isced_updated_count += 1

    wb.save(output_excel_path)
    wb.close()

    return {
        "total_rows": total_rows,
        "isced_updated_count": isced_updated_count,
        "dates_updated_count": dates_updated_count,
        "web_crawled_count": web_crawled_count,
        "output_path": str(output_excel_path),
    }
