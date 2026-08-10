"""Ana program pipeline v2 — Dolu Excel dosyalarındaki boş sütunları doldurur.

İş akışı:
1. Yüklenen dolu Excel dosyasındaki mevcut program satırlarını oku.
2. Opsiyonel olarak sunulan URL Kaynak Excel'indeki (v1'deki gibi) ve/veya 'Kabul Verisi Kaynağı' sütunundaki URL'leri birleştir.
3. 'Gerekli değil' yazan veya eksik olan URL'ler için resmî domain üzerinden hedefli URL keşfi yap.
4. Eksik (boş) kalan hücreleri tespit et.
5. Web sayfalarını indirip Genel Politika ve Program Özel veri çıkarıcılarını çalıştır.
6. 8 Sert QA Kuralı kontrolünü yap.
7. Orijinal dolu hücrelere HİÇ DOKUNMADAN sadece boş hücrelere yeni verileri yaz.
"""

from __future__ import annotations

import asyncio
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional, Any

from openpyxl import load_workbook

from src.config import get_settings
from src.domain_utils import extract_domain, is_safe_url
from src.fetch.page_fetcher import PageFetcher
from src.program_pipeline.checkpoint import CheckpointManager, PageCache
from src.program_pipeline.excel_writer_v2 import write_missing_fields_to_populated_excel, is_cell_empty
from src.program_pipeline.models import (
    ExtractedField, ProgramRecord, ProgramPipelineResult, SourceUniversityRow, TemplateInfo
)
from src.program_pipeline.normalizer import normalize_for_comparison
from src.program_pipeline.pipeline import ProgramPipeline
from src.program_pipeline.qa_exporter import evaluate_qa_report, save_qa_report
from src.program_pipeline.source_excel_reader import read_source_excel
from src.program_pipeline.template_reader import read_template_excel
from src.program_pipeline.targeted_search import run_targeted_search_for_program
from src.audit.evidence_store import EvidenceStore
from src.audit.review_queue import ReviewQueue

logger = logging.getLogger(__name__)


def extract_urls_from_source_cell(cell_value: Any) -> dict[str, str]:
    """'Kabul Verisi Kaynağı' hücresindeki etiketli veya ham URL'leri ayrıştır.

    Örnek hücre metni:
    "Program: https://www.academyart.edu/academics/
    Kabul: https://www.academyart.edu/admissions/international/
    Ücret: https://www.academyart.edu/admissions/cost-financial-aid/"
    """
    urls: dict[str, str] = {}
    if not cell_value:
        return urls

    text = str(cell_value)

    # Etiketli ayrıştırma
    prog_m = re.search(r"Program:\s*(https?://[^\s]+)", text, re.I)
    if prog_m:
        urls["program_url"] = prog_m.group(1).strip()

    adm_m = re.search(r"Kabul:\s*(https?://[^\s]+)", text, re.I)
    if adm_m:
        urls["admission_url"] = adm_m.group(1).strip()

    lang_m = re.search(r"Dil:\s*(https?://[^\s]+)", text, re.I)
    if lang_m:
        urls["language_url"] = lang_m.group(1).strip()

    tuit_m = re.search(r"Ücret:\s*(https?://[^\s]+)", text, re.I)
    if tuit_m:
        urls["tuition_url"] = tuit_m.group(1).strip()

    # Eğer etiket yoksa ama ham URL'ler varsa
    if not urls:
        raw_urls = re.findall(r"https?://[^\s\n,]+", text)
        if raw_urls:
            urls["program_url"] = raw_urls[0]
            if len(raw_urls) > 1:
                urls["admission_url"] = raw_urls[1]

    return urls


class ProgramPipelineV2:
    """Dolu Excel dosyasındaki eksik sütunları doldurma pipeline'ı (v2)."""

    def __init__(
        self,
        confidence_threshold: float = 0.85,
        resume: bool = True,
        progress_callback: Optional[Callable[[dict], None]] = None,
    ):
        self.confidence_threshold = confidence_threshold
        self.resume = resume
        self.progress_callback = progress_callback
        self.base_pipeline = ProgramPipeline(
            confidence_threshold=confidence_threshold,
            resume=resume,
            progress_callback=progress_callback,
        )
        self._log_lines: list[str] = []

    def _log(self, msg: str, *args) -> None:
        text = msg % args if args else msg
        logger.info(text)
        self._log_lines.append(f"[{datetime.now().strftime('%H:%M:%S')}] {text}")

    def _notify(
        self,
        uni_done: int,
        uni_total: int,
        prog_done: int,
        prog_total: int,
        uni_name: str = "",
    ) -> None:
        if self.progress_callback:
            try:
                self.progress_callback({
                    "uni_done": uni_done,
                    "uni_total": uni_total,
                    "prog_done": prog_done,
                    "prog_total": prog_total,
                    "uni_name": uni_name,
                    "log": self._log_lines[-1] if self._log_lines else "",
                })
            except Exception:
                pass

    async def _ensure_university_urls(self, row: SourceUniversityRow) -> SourceUniversityRow:
        """Eğer bir URL (Kabul, Dil, Ücret veya Program) eksikse veya 'gerekli değil' ise,
        üniversitenin resmî domaini üzerinden hedefli web aramasıyla URL keşfi yap.
        """
        domain = extract_domain(row.official_website or row.program_url or row.admission_url or "")
        if not domain:
            return row

        # 1. Kabul URL'si eksik mi veya 'gerekli değil' mi?
        if not row.admission_url or row.admission_not_needed:
            discovered = await run_targeted_search_for_program(domain, row.university_name, "admission requirements international requirements")
            if discovered:
                row.admission_url = discovered[0]
                row.admission_not_needed = False
                self._log("🔍 '%s' için Kabul URL'si keşfedildi: %s", row.university_name, discovered[0])

        # 2. Dil URL'si eksik mi veya 'gerekli değil' mi?
        if not row.language_url or row.language_not_needed:
            discovered = await run_targeted_search_for_program(domain, row.university_name, "english proficiency ielts toefl requirements")
            if discovered:
                row.language_url = discovered[0]
                row.language_not_needed = False
                self._log("🔍 '%s' için Dil URL'si keşfedildi: %s", row.university_name, discovered[0])

        # 3. Ücret URL'si eksik mi veya 'gerekli değil' mi?
        if not row.tuition_url or row.tuition_not_needed:
            discovered = await run_targeted_search_for_program(domain, row.university_name, "tuition and fees cost of attendance")
            if discovered:
                row.tuition_url = discovered[0]
                row.tuition_not_needed = False
                self._log("🔍 '%s' için Ücret URL'si keşfedildi: %s", row.university_name, discovered[0])

        # 4. Program URL'si eksik mi veya 'gerekli değil' mi?
        if not row.program_url or row.program_not_needed:
            discovered = await run_targeted_search_for_program(domain, row.university_name, "academic degree programs majors")
            if discovered:
                row.program_url = discovered[0]
                row.program_not_needed = False
                self._log("🔍 '%s' için Program URL'si keşfedildi: %s", row.university_name, discovered[0])

        return row

    async def run(
        self,
        populated_excel_path: Path | str,
        output_dir: Path | str,
        source_excel_path: Optional[Path | str] = None,
    ) -> ProgramPipelineResult:
        """Dolu Excel dosyasını oku ve sadece eksik hücreleri doldur.

        Args:
            populated_excel_path: Dolu program Excel dosyası
            output_dir: Çıktı klasörü
            source_excel_path: Opsiyonel v1 tarzı URL Kaynak Excel dosyası
        """
        settings = get_settings()
        populated_excel_path = Path(populated_excel_path)
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        template_info = read_template_excel(populated_excel_path)
        wb = load_workbook(populated_excel_path, data_only=True)
        sheet_name = template_info.sheet_name
        ws = wb[sheet_name]
        header_row = template_info.header_row

        column_map = template_info.column_map
        uni_col = column_map.get("Üniversite Adı") or column_map.get("University Name") or 2
        prog_col = column_map.get("Program Adı") or column_map.get("Program Name") or 3
        source_col = column_map.get("Kabul Verisi Kaynağı")

        # 1. Dolu Excel'deki üniversiteleri ve URL'lerini grupla
        uni_data: dict[str, dict[str, Any]] = {}

        for row_idx in range(header_row + 1, (ws.max_row or header_row) + 1):
            uni_val = ws.cell(row=row_idx, column=uni_col).value
            prog_val = ws.cell(row=row_idx, column=prog_col).value
            source_val = ws.cell(row=row_idx, column=source_col).value if source_col else None

            if not uni_val or not prog_val:
                continue

            uni_name = str(uni_val).strip()
            prog_name = str(prog_val).strip()

            if uni_name not in uni_data:
                parsed_urls = extract_urls_from_source_cell(source_val)
                uni_data[uni_name] = {
                    "university_name": uni_name,
                    "programs": [],
                    "urls": parsed_urls,
                }

            uni_data[uni_name]["programs"].append(prog_name)

        wb.close()

        # 2. Opsiyonel URL Kaynak Excel'i oku (v1 gibi)
        external_source_map: dict[str, SourceUniversityRow] = {}
        if source_excel_path and Path(source_excel_path).exists():
            self._log("URL Kaynak Excel'i okunuyor: %s", Path(source_excel_path).name)
            try:
                ext_rows = read_source_excel(source_excel_path)
                for r in ext_rows:
                    norm_n = normalize_for_comparison(r.university_name)
                    if norm_n:
                        external_source_map[norm_n] = r
                self._log("URL Kaynak Excel'inden %d üniversite URL verisi yüklendi.", len(ext_rows))
            except Exception as e:
                self._log("⚠️ URL Kaynak Excel okunurken uyarı: %s", e)

        # 3. SourceUniversityRow objelerini oluştur (Birleştirme mekanizması)
        source_rows: list[SourceUniversityRow] = []
        for uni_name, u_info in uni_data.items():
            cell_urls = u_info["urls"]

            # Dış URL Excel'inde bu üniversite var mı?
            norm_uni = normalize_for_comparison(uni_name)
            ext_row = external_source_map.get(norm_uni)

            prog_url = (ext_row.program_url if ext_row and ext_row.program_url else cell_urls.get("program_url", ""))
            adm_url = (ext_row.admission_url if ext_row and ext_row.admission_url else cell_urls.get("admission_url", ""))
            lang_url = (ext_row.language_url if ext_row and ext_row.language_url else cell_urls.get("language_url", ""))
            tuit_url = (ext_row.tuition_url if ext_row and ext_row.tuition_url else cell_urls.get("tuition_url", ""))
            off_site = (ext_row.official_website if ext_row and ext_row.official_website else prog_url)

            prog_nn = ext_row.program_not_needed if ext_row else False
            adm_nn = ext_row.admission_not_needed if ext_row else False
            lang_nn = ext_row.language_not_needed if ext_row else False
            tuit_nn = ext_row.tuition_not_needed if ext_row else False

            s_row = SourceUniversityRow(
                row_number=len(source_rows) + 1,
                university_name=uni_name,
                official_website=off_site,
                program_url=prog_url,
                admission_url=adm_url,
                language_url=lang_url,
                tuition_url=tuit_url,
                program_not_needed=prog_nn,
                admission_not_needed=adm_nn,
                language_not_needed=lang_nn,
                tuition_not_needed=tuit_nn,
            )
            source_rows.append(s_row)

        audit = EvidenceStore(settings.audit_dir)
        review_q = ReviewQueue(settings.audit_dir / "review_queue.jsonl")

        all_programs: list[ProgramRecord] = []
        rejected_candidates_summary: list[dict] = []

        uni_total = len(source_rows)

        # 4. Her üniversite için tarama ve çıkarma yap
        for uni_idx, row in enumerate(source_rows):
            uni_name = row.university_name
            self._log("⚡ v2 İşleniyor: %s (%d/%d)", uni_name, uni_idx + 1, uni_total)
            self._notify(uni_idx, uni_total, 0, len(uni_data[uni_name]["programs"]), uni_name)

            try:
                # 'Gerekli değil' yazan veya eksik olan URL'leri web aramasıyla keşfet
                row = await self._ensure_university_urls(row)

                programs, uni_stats, rej_list = await self.base_pipeline._process_university(
                    row=row,
                    template_info=template_info,
                    audit=audit,
                    review_q=review_q,
                    uni_idx=uni_idx,
                    uni_total=uni_total,
                )
                all_programs.extend(programs)
                rejected_candidates_summary.extend(rej_list)

            except Exception as e:
                self._log("HATA (%s): %s", uni_name, e)
                logger.exception("v2 Üniversite işleme hatası: %s", uni_name)

        # 5. QA Değerlendirmesi
        qa_report = evaluate_qa_report(
            programs=all_programs,
            total_unis=uni_total,
            rejected_candidates=rejected_candidates_summary,
        )
        qa_json_path = output_dir / "qa_report.json"
        save_qa_report(qa_report, qa_json_path)

        output_path = output_dir / f"v2_doldurulmus_{populated_excel_path.name}"

        write_stats = {
            "total_rows": 0,
            "existing_cells_preserved": 0,
            "empty_cells_filled": 0,
            "empty_cells_remaining": 0,
        }

        if all_programs:
            self._log("v2 Eksik sütunlar yazılıyor (%d program kaydı)...", len(all_programs))
            write_stats = write_missing_fields_to_populated_excel(
                populated_excel_path=populated_excel_path,
                output_path=output_path,
                program_records=all_programs,
                column_map=column_map,
                sheet_name=sheet_name,
                header_row=header_row,
                confidence_threshold=self.confidence_threshold,
            )

        return ProgramPipelineResult(
            output_path=str(output_path) if output_path.exists() else "",
            audit_path=str(settings.audit_dir / "audit.jsonl"),
            review_queue_path=str(settings.audit_dir / "review_queue.jsonl"),
            total_programs=len(all_programs),
            cells_written=write_stats["empty_cells_filled"],
            cells_empty=write_stats["empty_cells_remaining"],
            needs_review_count=0,
            errors=[],
            missing_rate=0.0,
        )
