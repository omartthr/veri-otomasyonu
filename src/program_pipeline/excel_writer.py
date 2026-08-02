"""Program pipeline Excel yazıcı — Programlar sheetine yazar.

Kurallar:
- Template workbook'u kopyala (orijinale dokunma)
- Programlar sheetini bul ve kullan
- Sütun eşleştirmesini başlık adına göre yap (sabit harf KULLANMA)
- Mevcut satırları silme; yeni satırlara template stilini kopyala
- Boş hücre = cell.value = None (asla None/N/A/null yazma)
- openpyxl comment olarak kaynak URL ekle
- Formüllere dokunma
- Atomik yazma: geçici dosyaya yaz, sonra taşı
"""

from __future__ import annotations

import logging
import shutil
import tempfile
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from src.program_pipeline.models import ExtractedField, ProgramRecord, TemplateInfo
from src.program_pipeline.program_quality import normalize_program_key

logger = logging.getLogger(__name__)

# Boş hücre olarak yazılmamalı
_EMPTY_SENTINEL = {None, "null", "none", "n/a", "na", "—", "-", ""}


def _cell_value(field: ExtractedField) -> Optional[object]:
    """Hücreye yazılacak değeri üret. Boşsa None döndür."""
    if field.value is None:
        return None
    val = field.value

    # Boolean değerleri metin yap
    if isinstance(val, bool):
        return "Evet" if val else "Hayır"

    # Boş sentinel
    if str(val).strip().lower() in _EMPTY_SENTINEL:
        return None

    return val


def _copy_row_style(
    ws: Worksheet,
    source_row: int,
    target_row: int,
    max_col: int,
) -> None:
    """Kaynak satırın hücre stilini hedef satıra kopyala."""
    from copy import copy
    for col in range(1, max_col + 1):
        src_cell = ws.cell(row=source_row, column=col)
        tgt_cell = ws.cell(row=target_row, column=col)
        if src_cell.has_style:
            tgt_cell.font = copy(src_cell.font)
            tgt_cell.border = copy(src_cell.border)
            tgt_cell.fill = copy(src_cell.fill)
            tgt_cell.number_format = src_cell.number_format
            tgt_cell.protection = copy(src_cell.protection)
            tgt_cell.alignment = copy(src_cell.alignment)


def _add_comment(ws: Worksheet, row: int, col: int, text: str) -> None:
    """Hücreye kaynak URL yorumu ekle."""
    try:
        comment = Comment(text[:500], "Üniversite Veri Botu")
        ws.cell(row=row, column=col).comment = comment
    except Exception:
        pass  # Yorum hatası kritik değil


def write_programs_to_excel(
    programs: list[ProgramRecord],
    template_path: Path | str,
    output_path: Path | str,
    template_info: TemplateInfo,
    confidence_threshold: float = 0.85,
) -> Path:
    """Program kayıtlarını template Excel'e yaz.

    Args:
        programs: Yazılacak program kayıtları
        template_path: Orijinal template dosyası
        output_path: Çıktı dosyası (geçici dosyaya yaz, sonra taşı)
        template_info: Programlar sheet meta bilgileri
        confidence_threshold: Bu değerin altındaki alanları boş bırak

    Returns:
        Yazılan dosyanın yolu

    Raises:
        FileNotFoundError: Template bulunamazsa
        ValueError: Programlar sheet bulunamazsa
    """
    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.exists():
        raise FileNotFoundError(f"Template bulunamadı: {template_path}")

    # Atomik yazma: önce geçici dosyaya
    with tempfile.NamedTemporaryFile(
        suffix=".xlsx", delete=False, dir=output_path.parent
    ) as tmp:
        tmp_path = Path(tmp.name)

    try:
        # Template'i geçici dosyaya kopyala
        shutil.copy2(template_path, tmp_path)

        # Workbook'u aç (formüller korunacak)
        wb = load_workbook(tmp_path, data_only=False)

        sheet_name = template_info.sheet_name
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"'{sheet_name}' sheet bulunamadı. Mevcut: {wb.sheetnames}"
            )

        ws = wb[sheet_name]
        header_row = template_info.header_row
        col_map = template_info.column_map  # başlık → sütun numarası

        # Stil kopyalamak için kullanılacak referans satır (başlık + 1)
        style_source_row = header_row + 1

        # Mevcut son satırı bul (yeni satırlar buradan eklenir)
        current_last_row = ws.max_row
        if current_last_row < header_row:
            current_last_row = header_row

        stats = {"written": 0, "empty": 0, "programs": 0}

        existing_keys = {
            (uni.strip().casefold(), normalize_program_key(prog))
            for uni, prog in template_info.existing_programs
        }

        for program in programs:
            if program.status != "accepted" or not program.program_name.value:
                continue
            dedupe_key = (
                program.university_name.strip().casefold(),
                normalize_program_key(str(program.program_name.value)),
            )
            if dedupe_key in existing_keys:
                logger.info(
                    "Mevcut program satırı atlandı: %s / %s",
                    program.university_name,
                    program.program_name.value,
                )
                continue
            existing_keys.add(dedupe_key)

            target_row = current_last_row + 1
            current_last_row = target_row
            stats["programs"] += 1

            # Stil kopyala
            if style_source_row <= ws.max_row:
                _copy_row_style(
                    ws,
                    source_row=style_source_row,
                    target_row=target_row,
                    max_col=ws.max_column or 1,
                )

            # Sabit alanlar: Üniversite ID, Üniversite Adı, Program Adı, Program URL
            _write_fixed_fields(ws, target_row, col_map, program, stats)

            # Değişken alanlar: ExtractedField'lar
            for field_name, field in program.fields.items():
                col = col_map.get(field_name)
                if col is None:
                    continue

                val = _cell_value(field)
                cell = ws.cell(row=target_row, column=col)

                if val is None or field.confidence < confidence_threshold:
                    cell.value = None
                    stats["empty"] += 1
                else:
                    cell.value = val
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    stats["written"] += 1

                    # Kaynak URL yorumu
                    if field.source_url:
                        comment_text = f"Kaynak: {field.source_url}"
                        if field.evidence:
                            comment_text += f"\nKanıt: {field.evidence[:200]}"
                        _add_comment(ws, target_row, col, comment_text)

        wb.save(tmp_path)
        wb.close()

        # Atomik taşıma
        shutil.move(str(tmp_path), str(output_path))

        logger.info(
            "Excel yazıldı: %s | program=%d, hücre=%d, boş=%d",
            output_path.name, stats["programs"], stats["written"], stats["empty"],
        )
        return output_path

    except Exception:
        # Temizlik
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise


def _write_fixed_fields(
    ws: Worksheet,
    row: int,
    col_map: dict[str, int],
    program: ProgramRecord,
    stats: dict,
) -> None:
    """Sabit alanları yaz: ID, isim, program adı, URL."""
    def _write(col_name: str, value: object) -> None:
        col = col_map.get(col_name)
        if col and value is not None:
            ws.cell(row=row, column=col, value=value)
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
            stats["written"] += 1

    _write("Üniversite ID", program.university_id or None)
    _write("Üniversite Adı", program.university_name or None)
    _write(
        "Program Adı",
        program.program_name.value if program.program_name.value is not None else None,
    )
    program_name_col = col_map.get("Program Adı")
    if program_name_col and program.program_name.source_url:
        _add_comment(
            ws,
            row,
            program_name_col,
            f"Program kaynağı: {program.program_name.source_url}",
        )

    # Program URL (varsa)
    if program.program_url:
        # Program sayfasını kabul kaynağı diye yazmak semantik olarak yanlıştı.
        # Yalnızca template'te ayrı bir Program URL sütunu varsa buraya yaz.
        col = col_map.get("Program URL")
        if col:
            ws.cell(row=row, column=col, value=program.program_url)


def get_output_filename(
    template_path: Path | str,
    output_dir: Path | str,
) -> Path:
    """Çıktı dosya adını oluştur."""
    template_path = Path(template_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = template_path.stem
    base = output_dir / f"{stem}_programlar_doldurulmus.xlsx"

    if not base.exists():
        return base

    for i in range(2, 10_000):
        candidate = output_dir / f"{stem}_programlar_doldurulmus_{i}.xlsx"
        if not candidate.exists():
            return candidate

    raise FileExistsError("Çıktı dosyası oluşturulamadı.")
