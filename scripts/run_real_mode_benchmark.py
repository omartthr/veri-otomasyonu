from __future__ import annotations

import asyncio
import json
import shutil
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.models import DecisionType, UniversityRow
from src.pipeline import ResearchPipeline
from src.program_pipeline.pipeline import ProgramPipeline
from src.program_pipeline.pipeline_v2 import ProgramPipelineV2
from src.program_pipeline.source_excel_reader import read_source_excel
from src.program_pipeline.template_reader import read_template_excel
from src.program_pipeline.university_matcher import match_universities


INPUT = ROOT / "data" / "input"
RESULTS = ROOT / "test-results" / "real-benchmark"
SOURCE = INPUT / "source_united_states_of_america_mississippi_program_kabul_takip (1).xlsx"
TEMPLATE = INPUT / "template_united_states_of_america_mississippi_university_country_review.xlsx"
POPULATED = INPUT / "v2_populated_united_states_of_america_arkansas_university_country_review (1).xlsx"
V2_SOURCE = INPUT / "v2_source_united_states_of_america_arkansas_program_kabul_takip (2).xlsx"


def timed_result(started: float, **values):
    return {"duration_seconds": round(time.perf_counter() - started, 2), **values}


def workbook_fill(path: Path) -> dict:
    info = read_template_excel(path)
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[info.sheet_name]
    values = []
    for row in sheet.iter_rows(min_row=info.header_row + 1):
        if any(cell.value not in (None, "") for cell in row):
            values.extend(cell.value for cell in row)
    workbook.close()
    total = len(values)
    filled = sum(value not in (None, "") for value in values)
    return {"cells": total, "filled": filled, "rate": round(filled / total, 4) if total else 0.0}


def make_single_university_populated() -> tuple[Path, Path, str]:
    RESULTS.mkdir(parents=True, exist_ok=True)
    populated_copy = RESULTS / "v2-single-university-input.xlsx"
    source_copy = RESULTS / "v2-single-university-source.xlsx"
    shutil.copy2(POPULATED, populated_copy)
    shutil.copy2(V2_SOURCE, source_copy)

    info = read_template_excel(populated_copy)
    workbook = load_workbook(populated_copy)
    sheet = workbook[info.sheet_name]
    uni_col = info.column_map.get("Üniversite Adı") or info.column_map.get("University Name") or 2
    first_university = None
    keep_rows = []
    for row_index in range(info.header_row + 1, sheet.max_row + 1):
        value = sheet.cell(row_index, uni_col).value
        if value and first_university is None:
            first_university = str(value).strip()
        if value and str(value).strip() == first_university:
            keep_rows.append(row_index)
    for row_index in range(sheet.max_row, info.header_row, -1):
        if row_index not in keep_rows:
            sheet.delete_rows(row_index)
    workbook.save(populated_copy)
    workbook.close()

    rows = read_source_excel(source_copy)
    matching = [row for row in rows if row.university_name == first_university]
    if not matching:
        raise RuntimeError(f"v2 source row not found: {first_university}")
    workbook = load_workbook(source_copy)
    sheet = workbook[workbook.sheetnames[0]]
    for row_index in range(sheet.max_row, 1, -1):
        name = sheet.cell(row_index, 2).value
        if name and str(name).strip() != first_university:
            sheet.delete_rows(row_index)
    workbook.save(source_copy)
    workbook.close()
    return populated_copy, source_copy, first_university


async def main() -> None:
    RESULTS.mkdir(parents=True, exist_ok=True)
    report = {"sample_strategy": "one real university per mode", "modes": {}}

    source_row = read_source_excel(SOURCE)[0]
    url_row = UniversityRow(
        row_number=source_row.row_number,
        university_id=source_row.university_id,
        university_name=source_row.university_name,
        website=source_row.official_website,
        program_missing="Evet",
        admission_missing="Evet",
        language_missing="Evet",
        tuition_missing="Evet",
    )
    started = time.perf_counter()
    url_result = (await ResearchPipeline().run([url_row]))[0]
    accepted = sum(
        result is not None and result.decision == DecisionType.ACCEPTED
        for result in url_result.best_results.values()
    )
    report["modes"]["url_discovery"] = timed_result(
        started,
        university=source_row.university_name,
        requested=4,
        accepted=accepted,
        completion_rate=round(accepted / 4, 4),
        status=url_result.status.value,
        results={
            category.value: {
                "url": result.candidate.url if result else None,
                "decision": result.decision.value if result else "not_found",
                "score": result.score.total if result else 0,
            }
            for category, result in url_result.best_results.items()
        },
        notes=url_result.notes,
    )

    template_info = read_template_excel(TEMPLATE)
    matches = match_universities([source_row], template_info.university_names)
    started = time.perf_counter()
    v1_result = await ProgramPipeline(confidence_threshold=0.85, resume=False).run(
        source_rows=[source_row],
        template_info=template_info,
        matches=matches,
        template_path=TEMPLATE,
        output_dir=RESULTS / "v1",
    )
    v1_qa_path = RESULTS / "v1" / "qa_report.json"
    v1_qa = json.loads(v1_qa_path.read_text(encoding="utf-8")) if v1_qa_path.exists() else {}
    v1_total = v1_result.cells_written + v1_result.cells_empty
    report["modes"]["program_v1"] = timed_result(
        started,
        university=source_row.university_name,
        programs=v1_result.total_programs,
        cells_written=v1_result.cells_written,
        cells_empty=v1_result.cells_empty,
        completion_rate=round(v1_result.cells_written / v1_total, 4) if v1_total else 0.0,
        needs_review=v1_result.needs_review_count,
        errors=v1_result.errors,
        output_path=v1_result.output_path,
        qa=v1_qa,
    )

    populated_copy, v2_source_copy, v2_university = make_single_university_populated()
    before = workbook_fill(populated_copy)
    started = time.perf_counter()
    v2_result = await ProgramPipelineV2(confidence_threshold=0.85, resume=False).run(
        populated_excel_path=populated_copy,
        output_dir=RESULTS / "v2",
        source_excel_path=v2_source_copy,
    )
    output_path = Path(v2_result.output_path) if v2_result.output_path else None
    after = workbook_fill(output_path) if output_path and output_path.exists() else before
    report["modes"]["program_v2"] = timed_result(
        started,
        university=v2_university,
        programs=v2_result.total_programs,
        cells_written=v2_result.cells_written,
        cells_empty=v2_result.cells_empty,
        completion_rate=round(v2_result.cells_written / (v2_result.cells_written + v2_result.cells_empty), 4)
        if v2_result.cells_written + v2_result.cells_empty else 0.0,
        workbook_fill_before=before,
        workbook_fill_after=after,
        existing_cells_preserved=after["filled"] >= before["filled"],
        errors=v2_result.errors,
        output_path=v2_result.output_path,
    )

    combined_success = accepted + (1 if v1_result.total_programs > 0 else 0)
    report["modes"]["combined"] = {
        "university": source_row.university_name,
        "duration_seconds": round(
            report["modes"]["url_discovery"]["duration_seconds"]
            + report["modes"]["program_v1"]["duration_seconds"],
            2,
        ),
        "stages_completed": combined_success,
        "stages_total": 5,
        "completion_rate": round(combined_success / 5, 4),
        "note": "URL discovery followed by the same v1 automation run.",
    }

    report_path = RESULTS / "benchmark-results.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    asyncio.run(main())
